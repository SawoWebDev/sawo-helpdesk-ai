"""Combined FAQ + Library export/import: one workbook, two sheets, so the
whole knowledge base can be backed up or moved to another deployment without
recrawling Library sources or re-entering FAQs by hand."""

import io
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.crud.category import build_category_path_map, get_or_create_category_path
from app.crud.faq import create_faq, list_faqs_all
from app.crud.vault import create_vault_entry, list_vault_entries_all
from app.rag.reindex import embed_entry
from app.rag.vault_reindex import embed_vault_entry

FAQ_SHEET_NAME = "FAQs"
LIBRARY_SHEET_NAME = "Library"

FAQ_HEADERS = ["Category", "Question", "Answer", "Image URL", "Reference URL"]
FAQ_REQUIRED_HEADERS = ["Category", "Question", "Answer"]

LIBRARY_HEADERS = ["Category", "Title", "Content", "Tags", "Source URL", "Memory Enabled"]
LIBRARY_REQUIRED_HEADERS = ["Category", "Title", "Content"]


@dataclass
class RowError:
    row_number: int
    reason: str


@dataclass
class ImportSummary:
    created: int = 0
    skipped: int = 0
    failed: int = 0
    errors: list[RowError] = field(default_factory=list)


@dataclass
class KnowledgeBaseImportSummary:
    faqs: ImportSummary
    library: ImportSummary


async def export_knowledge_base(db: AsyncSession) -> bytes:
    """Same column layout each entity's own standalone export already uses
    (see services/excel_import.py), just as two sheets in one file."""
    faq_entries = await list_faqs_all(db)
    vault_entries = await list_vault_entries_all(db)
    category_paths = await build_category_path_map(db)

    wb = Workbook()
    faq_ws = wb.active
    faq_ws.title = FAQ_SHEET_NAME
    faq_ws.append(FAQ_HEADERS)
    for entry in faq_entries:
        category_path = category_paths.get(entry.category_id, "") if entry.category_id else ""
        faq_ws.append(
            [
                category_path,
                entry.question,
                entry.answer,
                (entry.image_urls or [""])[0],
                (entry.reference_urls or [""])[0],
            ]
        )

    library_ws = wb.create_sheet(LIBRARY_SHEET_NAME)
    library_ws.append(LIBRARY_HEADERS)
    for entry in vault_entries:
        category_path = category_paths.get(entry.category_id, "") if entry.category_id else ""
        library_ws.append(
            [
                category_path,
                entry.title,
                entry.content,
                ", ".join(entry.tags or []),
                entry.source_url or "",
                "TRUE" if entry.memory_enabled else "FALSE",
            ]
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _row_cells(headers: list[str], row: tuple) -> dict[str, str]:
    col_index = {h: i for i, h in enumerate(headers)}

    def cell(name: str) -> str:
        idx = col_index.get(name)
        if idx is None or idx >= len(row):
            return ""
        value = row[idx]
        return str(value).strip() if value is not None else ""

    return {name: cell(name) for name in headers}


async def _import_faq_sheet(db: AsyncSession, ws) -> ImportSummary:
    summary = ImportSummary()
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return summary

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    missing = [h for h in FAQ_REQUIRED_HEADERS if h not in headers]
    if missing:
        summary.failed += 1
        summary.errors.append(RowError(row_number=1, reason=f"Missing required column(s): {', '.join(missing)}"))
        return summary

    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None for cell in row):
            continue
        cells = _row_cells(headers, row)
        category_path = cells.get("Category", "")
        question = cells.get("Question", "")
        answer = cells.get("Answer", "")
        image_url = cells.get("Image URL", "")
        reference_url = cells.get("Reference URL", "")

        missing_fields = [
            name
            for name, value in [("Category", category_path), ("Question", question), ("Answer", answer)]
            if not value
        ]
        if missing_fields:
            summary.skipped += 1
            summary.errors.append(
                RowError(row_number=row_number, reason=f"Missing required field(s): {', '.join(missing_fields)}")
            )
            continue

        try:
            category = await get_or_create_category_path(db, category_path)
            entry = await create_faq(
                db,
                question=question,
                answer=answer,
                category_id=category.id,
                image_urls=[image_url] if image_url else [],
                reference_urls=[reference_url] if reference_url else [],
                source="import",
            )
            await embed_entry(db, entry)
            await db.commit()
            summary.created += 1
        except Exception as exc:
            await db.rollback()
            summary.failed += 1
            summary.errors.append(RowError(row_number=row_number, reason=str(exc)))

    return summary


async def _import_library_sheet(db: AsyncSession, ws) -> ImportSummary:
    summary = ImportSummary()
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return summary

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    missing = [h for h in LIBRARY_REQUIRED_HEADERS if h not in headers]
    if missing:
        summary.failed += 1
        summary.errors.append(RowError(row_number=1, reason=f"Missing required column(s): {', '.join(missing)}"))
        return summary

    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None for cell in row):
            continue
        cells = _row_cells(headers, row)
        category_path = cells.get("Category", "")
        title = cells.get("Title", "")
        content = cells.get("Content", "")
        tags_raw = cells.get("Tags", "")
        source_url = cells.get("Source URL", "")
        memory_raw = cells.get("Memory Enabled", "")

        missing_fields = [
            name for name, value in [("Category", category_path), ("Title", title), ("Content", content)] if not value
        ]
        if missing_fields:
            summary.skipped += 1
            summary.errors.append(
                RowError(row_number=row_number, reason=f"Missing required field(s): {', '.join(missing_fields)}")
            )
            continue

        try:
            category = await get_or_create_category_path(db, category_path)
            tags = [t.strip() for t in tags_raw.split(",") if t.strip()]
            # Default to memory-enabled (searchable in chat) when the column is
            # blank, since the point of importing is to make this content
            # answerable again without recrawling — not to re-stage it as an
            # inert copy an admin has to remember to switch on.
            memory_enabled = memory_raw.strip().upper() != "FALSE" if memory_raw else True
            entry = await create_vault_entry(
                db,
                title=title,
                content=content,
                category_id=category.id,
                tags=tags,
                source_type="import",
                source_url=source_url or None,
            )
            entry.memory_enabled = memory_enabled
            if memory_enabled:
                await embed_vault_entry(db, entry)
            await db.commit()
            summary.created += 1
        except Exception as exc:
            await db.rollback()
            summary.failed += 1
            summary.errors.append(RowError(row_number=row_number, reason=str(exc)))

    return summary


async def import_knowledge_base(db: AsyncSession, file_bytes: bytes) -> KnowledgeBaseImportSummary:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        error = ImportSummary(failed=1, errors=[RowError(row_number=0, reason=f"Could not read workbook: {exc}")])
        return KnowledgeBaseImportSummary(faqs=error, library=ImportSummary())

    faq_summary = (
        await _import_faq_sheet(db, wb[FAQ_SHEET_NAME]) if FAQ_SHEET_NAME in wb.sheetnames else ImportSummary()
    )
    library_summary = (
        await _import_library_sheet(db, wb[LIBRARY_SHEET_NAME])
        if LIBRARY_SHEET_NAME in wb.sheetnames
        else ImportSummary()
    )
    return KnowledgeBaseImportSummary(faqs=faq_summary, library=library_summary)
