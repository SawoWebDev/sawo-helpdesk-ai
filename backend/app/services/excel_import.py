import io
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.crud.category import build_category_path_map, get_or_create_category_path
from app.crud.faq import create_faq, list_faqs_all
from app.rag.reindex import embed_entry

TEMPLATE_HEADERS = ["Category", "Question", "Answer", "Image URL", "Reference URL"]
REQUIRED_HEADERS = ["Category", "Question", "Answer"]


def generate_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FAQ Import"
    ws.append(TEMPLATE_HEADERS)
    ws.append(
        [
            "Billing",
            "How do I update my payment method?",
            "Go to Account Settings > Billing and click 'Update Payment Method'.",
            "",
            "https://example.com/billing-help",
        ]
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def export_faqs(
    db: AsyncSession,
    category_id: int | None = None,
    search: str | None = None,
) -> bytes:
    """Export FAQ entries (optionally filtered) using the same column layout as
    the import template, so the file can be edited and re-imported."""
    entries = await list_faqs_all(db, category_id=category_id, search=search)
    category_paths = await build_category_path_map(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "FAQ Export"
    ws.append(TEMPLATE_HEADERS)

    for entry in entries:
        category_path = category_paths.get(entry.category_id, "") if entry.category_id else ""
        ws.append(
            [
                category_path,
                entry.question,
                entry.answer,
                (entry.image_urls or [""])[0],
                (entry.reference_urls or [""])[0],
            ]
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class RowError:
    row_number: int
    reason: str


@dataclass
class ImportSummary:
    created: int = 0
    skipped: int = 0
    failed: int = 0
    errors: list[RowError] = field(default_factory=list)


async def import_workbook(db: AsyncSession, file_bytes: bytes) -> ImportSummary:
    summary = ImportSummary()

    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        summary.errors.append(RowError(row_number=0, reason=f"Could not read workbook: {exc}"))
        summary.failed += 1
        return summary

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        summary.errors.append(RowError(row_number=0, reason="Workbook is empty"))
        summary.failed += 1
        return summary

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        summary.errors.append(
            RowError(row_number=1, reason=f"Missing required column(s): {', '.join(missing)}")
        )
        summary.failed += 1
        return summary

    col_index = {h: i for i, h in enumerate(headers)}

    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None for cell in row):
            continue

        def cell(name: str) -> str:
            idx = col_index.get(name)
            if idx is None or idx >= len(row):
                return ""
            value = row[idx]
            return str(value).strip() if value is not None else ""

        category_path = cell("Category")
        question = cell("Question")
        answer = cell("Answer")
        image_url = cell("Image URL")
        reference_url = cell("Reference URL")

        missing_fields = []
        if not category_path:
            missing_fields.append("Category")
        if not question:
            missing_fields.append("Question")
        if not answer:
            missing_fields.append("Answer")

        if missing_fields:
            summary.skipped += 1
            summary.errors.append(
                RowError(
                    row_number=row_number,
                    reason=f"Missing required field(s): {', '.join(missing_fields)}",
                )
            )
            continue

        try:
            category = await get_or_create_category_path(db, category_path)
            image_urls = [image_url] if image_url else []
            reference_urls = [reference_url] if reference_url else []
            entry = await create_faq(
                db,
                question=question,
                answer=answer,
                category_id=category.id,
                image_urls=image_urls,
                reference_urls=reference_urls,
                source="import",
            )
            await embed_entry(db, entry)
            await db.commit()
            summary.created += 1
        except Exception as exc:
            await db.rollback()
            summary.failed += 1
            summary.errors.append(RowError(row_number=row_number, reason=str(exc)))

    return summary
