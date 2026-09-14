"""Text extraction for Library ingestion: PDF, DOCX, and crawled web pages.
Each parser returns plain text only — chunking happens separately in
app.services.chunking."""

import io

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

CRAWL_TIMEOUT = httpx.Timeout(30.0, connect=10.0)
CRAWL_USER_AGENT = "Mozilla/5.0 (compatible; HelpdeskLibraryBot/1.0)"


class ParseError(Exception):
    pass


def extract_pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ParseError(f"Failed to parse PDF: {exc}") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip())


def extract_docx_text(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:
        raise ParseError(f"Failed to parse DOCX: {exc}") from exc
    return "\n".join(p.strip() for p in paragraphs if p.strip())


def extract_xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ParseError(f"Failed to parse XLSX: {exc}") from exc

    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# Strip only what genuinely cannot render as readable text (raw script/style
# source code, SVG markup) — capture everything else verbatim, including nav
# menus, headers/footers, and forms. No guessing at "chrome" vs "real
# content": whatever text is actually on the crawled page goes into the
# Library as-is.
_STRIP_TAGS = [
    "script",
    "style",
    "svg",
]


def extract_html_text(html: str) -> str:
    soup = BeautifulSoup(html, "lxml")

    for tag_name in _STRIP_TAGS:
        for tag in soup.find_all(tag_name):
            tag.decompose()

    body = soup.body or soup
    text = body.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


async def fetch_url(url: str) -> tuple[str, str]:
    """Returns (html, final_url)."""
    try:
        async with httpx.AsyncClient(
            timeout=CRAWL_TIMEOUT, follow_redirects=True, headers={"User-Agent": CRAWL_USER_AGENT}
        ) as client:
            resp = await client.get(url)
            resp.raise_for_status()
            return resp.text, str(resp.url)
    except httpx.HTTPError as exc:
        raise ParseError(f"Failed to fetch URL: {exc}") from exc


SITEMAP_PATHS = ("/sitemap.xml", "/sitemap_index.xml", "/wp-sitemap.xml")
MAX_SITEMAP_DOCS = 50  # cap how many nested sitemap files a sitemap index can point to
MAX_SITEMAP_URLS = 2000  # cap total page URLs collected, so a huge site can't hang ingestion


async def _fetch_xml(client: httpx.AsyncClient, url: str) -> str | None:
    try:
        resp = await client.get(url)
        if resp.status_code != 200:
            return None
        return resp.text
    except httpx.HTTPError:
        return None


def _parse_sitemap_xml(xml_text: str) -> tuple[list[str], list[str]]:
    """Returns (page_urls, nested_sitemap_urls). A sitemap document is either
    a <urlset> of actual pages or a <sitemapindex> pointing to other sitemap
    documents — never both, per the sitemaps.org spec."""
    soup = BeautifulSoup(xml_text, "xml")

    nested = [loc.get_text(strip=True) for loc in soup.select("sitemapindex > sitemap > loc")]
    if nested:
        return [], nested

    # Only top-level <url><loc> (the page's own URL), not <image:image><image:loc>
    # nested inside it, which would otherwise get treated as a "page".
    pages = [url_tag.find("loc").get_text(strip=True) for url_tag in soup.find_all("url") if url_tag.find("loc")]
    return pages, []


async def discover_sitemap_urls(site_url: str) -> list[str]:
    """Given any URL on a site, finds and fully expands that site's sitemap
    (following sitemap-index nesting, common with WordPress/Yoast SEO) and
    returns every page URL it lists. Raises ParseError if no sitemap could be
    found at any of the standard locations."""
    parsed = httpx.URL(site_url)
    origin = f"{parsed.scheme}://{parsed.host}"

    async with httpx.AsyncClient(
        timeout=CRAWL_TIMEOUT, follow_redirects=True, headers={"User-Agent": CRAWL_USER_AGENT}
    ) as client:
        root_xml = None
        for path in SITEMAP_PATHS:
            root_xml = await _fetch_xml(client, origin + path)
            if root_xml:
                break
        if not root_xml:
            raise ParseError(
                f"No sitemap found at {origin} (tried {', '.join(SITEMAP_PATHS)}). "
                "Enter individual page URLs instead."
            )

        pages, nested = _parse_sitemap_xml(root_xml)
        all_pages: list[str] = list(pages)
        visited_docs = 1

        queue = list(nested)
        while queue and visited_docs < MAX_SITEMAP_DOCS and len(all_pages) < MAX_SITEMAP_URLS:
            doc_url = queue.pop(0)
            xml_text = await _fetch_xml(client, doc_url)
            visited_docs += 1
            if not xml_text:
                continue
            sub_pages, sub_nested = _parse_sitemap_xml(xml_text)
            all_pages.extend(sub_pages)
            queue.extend(sub_nested)

    # De-duplicate while preserving discovery order, and cap the final count.
    seen: set[str] = set()
    ordered_unique: list[str] = []
    for url in all_pages:
        if url not in seen:
            seen.add(url)
            ordered_unique.append(url)
    return ordered_unique[:MAX_SITEMAP_URLS]


def extract_text_for_file(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        return extract_pdf_text(data)
    if ext == "docx":
        return extract_docx_text(data)
    if ext == "xlsx":
        return extract_xlsx_text(data)
    raise ParseError(f"Unsupported file type: .{ext}")
