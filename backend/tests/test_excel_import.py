from openpyxl import Workbook, load_workbook

from app.services.excel_import import TEMPLATE_HEADERS, generate_template


def test_generate_template_headers():
    content = generate_template()
    import io

    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == TEMPLATE_HEADERS
    assert ws.max_row == 2  # header + one example row
