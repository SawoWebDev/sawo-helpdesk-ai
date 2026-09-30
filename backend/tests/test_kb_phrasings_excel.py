"""Alternate phrasings in the knowledge-base workbook (services/kb_export.py):
round trip to another deployment, repeated imports, bad references, and
workbooks exported before the Phrasings sheet existed."""

import io

import pytest_asyncio
from openpyxl import Workbook, load_workbook
from sqlalchemy import event, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.crud.faq_phrasing import create_phrasing
from app.db.base import Base, _setup_connection
from app.db.vec_store import (
    FAQ_PHRASING_VEC_TABLE,
    FAQ_QUESTION_VEC_TABLE,
    FAQ_VEC_TABLE,
    VAULT_VEC_TABLE,
    create_vec_table_sql,
)
from app.models.category import Category
from app.models.faq import FAQEntry
from app.models.faq_phrasing import FAQPhrasing
from app.models.user import User
from app.rag.saved_answers import find_exact_answer
from app.services.kb_export import (
    FAQ_HEADERS,
    FAQ_SHEET_NAME,
    LIBRARY_HEADERS,
    LIBRARY_SHEET_NAME,
    PHRASING_HEADERS,
    PHRASING_SHEET_NAME,
    export_knowledge_base,
    import_knowledge_base,
)
from tests.conftest import add_faq

E1_Q = "Innova Power Controller Error 1 (E1) — Temperature Sensor 1 (TS1) is not connected: what causes it and how do we fix it?"
E1_A = "E1 means Temperature Sensor 1 (TS1) is not connected. Check the TS1 wiring."
E4_Q = "Innova Power Controller Error 4 (E4) — Temperature Sensor 2 (TS2) is not connected: what causes it and how do we fix it?"
E4_A = "E4 means Temperature Sensor 2 (TS2) is not connected. Check the TS2 wiring."
P_ON = "Innova display shows E1, what is wrong?"
P_OFF = "What does E1 mean on my Innova controller?"
P_NEW = "Innova E1 error, how do I fix it?"


@pytest_asyncio.fixture
async def other_db(tmp_path):
    """A second, empty deployment to import into."""
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'other.db'}")
    event.listen(engine.sync_engine, "connect", lambda conn, _rec: conn.run_async(_setup_connection))
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
        for table in (FAQ_VEC_TABLE, VAULT_VEC_TABLE, FAQ_QUESTION_VEC_TABLE, FAQ_PHRASING_VEC_TABLE):
            await conn.execute(text(create_vec_table_sql(table)))
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        yield session
    await engine.dispose()


async def seed(db):
    category = Category(name="Innova")
    db.add(category)
    await db.commit()
    faq = await add_faq(db, E1_Q, E1_A)
    faq.category_id = category.id
    await db.commit()
    await create_phrasing(db, faq, P_ON, created_by_id=None)
    await create_phrasing(db, faq, P_OFF, created_by_id=None, enabled=False)
    return faq


def workbook(faq_rows=(), phrasing_rows=(), phrasing_headers=PHRASING_HEADERS, with_phrasings=True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = FAQ_SHEET_NAME
    ws.append(FAQ_HEADERS)
    for row in faq_rows:
        ws.append(row)
    wb.create_sheet(LIBRARY_SHEET_NAME).append(LIBRARY_HEADERS)
    if with_phrasings:
        sheet = wb.create_sheet(PHRASING_SHEET_NAME)
        sheet.append(list(phrasing_headers))
        for row in phrasing_rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def phrasings(db) -> dict[str, tuple[int, bool]]:
    rows = (await db.execute(select(FAQPhrasing))).scalars().all()
    return {p.phrasing: (p.faq_id, p.enabled) for p in rows}


# --- export ------------------------------------------------------------------------------


async def test_export_ties_each_phrasing_to_its_faq_question(db, embedder):
    await seed(db)
    wb = load_workbook(io.BytesIO(await export_knowledge_base(db)))
    rows = list(wb[PHRASING_SHEET_NAME].iter_rows(values_only=True))
    assert rows == [tuple(PHRASING_HEADERS), (E1_Q, P_ON, "TRUE"), (E1_Q, P_OFF, "FALSE")]
    assert [r[1] for r in wb[FAQ_SHEET_NAME].iter_rows(values_only=True)][1:] == [E1_Q]


# --- import ------------------------------------------------------------------------------


async def test_round_trip_to_another_deployment(db, other_db, embedder):
    await seed(db)
    summary = await import_knowledge_base(other_db, await export_knowledge_base(db))

    assert (summary.faqs.created, summary.phrasings.created, summary.phrasings.failed) == (1, 2, 0)
    [faq] = (await other_db.execute(select(FAQEntry))).scalars().all()
    assert await phrasings(other_db) == {P_ON: (faq.id, True), P_OFF: (faq.id, False)}
    assert (await find_exact_answer(other_db, P_ON)).hit  # served in the new deployment
    assert not (await find_exact_answer(other_db, P_OFF)).hit  # still disabled there


async def test_importing_the_same_file_again_creates_no_duplicates(db, other_db, embedder):
    await seed(db)
    exported = await export_knowledge_base(db)
    await import_knowledge_base(other_db, exported)
    again = await import_knowledge_base(other_db, exported)

    assert (again.phrasings.created, again.phrasings.unchanged, again.phrasings.failed) == (0, 2, 0)
    assert len(await phrasings(other_db)) == 2


async def test_existing_phrasings_are_kept_when_the_file_does_not_list_them(db, embedder):
    faq = await seed(db)
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, P_NEW, "TRUE")]))
    assert summary.phrasings.created == 1
    assert await phrasings(db) == {P_ON: (faq.id, True), P_OFF: (faq.id, False), P_NEW: (faq.id, True)}


async def test_reimport_never_changes_an_existing_phrasings_state(db, embedder):
    faq = await seed(db)
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, P_OFF, "TRUE")]))
    assert summary.phrasings.unchanged == 1
    assert (await phrasings(db))[P_OFF] == (faq.id, False)  # staff disabled it; the file doesn't override that


async def test_phrasing_for_a_missing_faq_is_refused(db, embedder):
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, P_NEW, "TRUE")]))
    assert summary.phrasings.failed == 1 and "No FAQ has this question" in summary.phrasings.errors[0].reason
    assert await phrasings(db) == {}


async def test_phrasing_already_mapped_to_another_faq_is_refused(db, embedder):
    await seed(db)
    await add_faq(db, E4_Q, E4_A)
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E4_Q, P_ON, "TRUE")]))
    assert summary.phrasings.failed == 1 and "already mapped to FAQ" in summary.phrasings.errors[0].reason


async def test_the_staff_ui_identifier_rules_apply(db, embedder):
    await add_faq(db, E1_Q, E1_A)
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, "Innova display shows E4, what is wrong?", "TRUE")]))
    assert summary.phrasings.failed == 1 and "doesn't match this FAQ's question" in summary.phrasings.errors[0].reason
    assert await phrasings(db) == {}


async def test_a_phrasing_can_be_its_faq_on_this_deployment_only_once(db, embedder):
    await add_faq(db, E1_Q, E1_A)
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, E1_Q.upper(), "TRUE")]))
    assert summary.phrasings.failed == 1 and "FAQ's own question" in summary.phrasings.errors[0].reason


async def test_rows_without_an_enabled_value_are_not_served_until_reviewed(db, embedder):
    await add_faq(db, E1_Q, E1_A)
    summary = await import_knowledge_base(
        db, workbook(phrasing_rows=[(E1_Q, P_NEW)], phrasing_headers=["FAQ Question", "Phrasing"])
    )
    assert summary.phrasings.created == 1
    assert (await phrasings(db))[P_NEW][1] is False
    assert not (await find_exact_answer(db, P_NEW)).hit


async def test_invalid_rows_are_reported(db, embedder):
    await add_faq(db, E1_Q, E1_A)
    summary = await import_knowledge_base(
        db, workbook(phrasing_rows=[(E1_Q, "", "TRUE"), (E1_Q, P_NEW, "maybe")])
    )
    assert summary.phrasings.skipped == 2 and summary.phrasings.created == 0
    missing_column = await import_knowledge_base(db, workbook(phrasing_headers=["Phrasing"]))
    assert "Missing required column" in missing_column.phrasings.errors[0].reason


async def test_draft_faqs_take_phrasings_as_in_the_staff_ui_but_are_not_served(db, embedder):
    await add_faq(db, E1_Q, E1_A, status="draft")
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, P_NEW, "TRUE")]))
    assert summary.phrasings.created == 1
    assert not (await find_exact_answer(db, P_NEW)).hit


async def test_importer_is_recorded_as_the_creator(db, embedder):
    admin = User(username="admin", password_hash="x", role="admin")
    db.add(admin)
    await db.commit()
    await add_faq(db, E1_Q, E1_A)
    await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, P_NEW, "TRUE")]), created_by_id=admin.id)
    [row] = (await db.execute(select(FAQPhrasing))).scalars().all()
    assert row.created_by_id == admin.id


async def test_workbook_without_a_phrasings_sheet_imports_as_before(db, embedder):
    summary = await import_knowledge_base(db, workbook(faq_rows=[["Heaters", E1_Q, E1_A, "", ""]], with_phrasings=False))
    assert summary.faqs.created == 1 and summary.faqs.failed == 0
    p = summary.phrasings
    assert (p.created, p.skipped, p.failed, p.unchanged, p.errors) == (0, 0, 0, 0, [])


async def test_duplicated_faq_rows_are_reported_not_guessed(db, embedder):
    await add_faq(db, E1_Q, E1_A)
    await add_faq(db, E1_Q, E1_A)  # the FAQs sheet has no duplicate check of its own
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[(E1_Q, P_NEW, "TRUE")]))
    assert summary.phrasings.failed == 1 and "Several FAQs have this question" in summary.phrasings.errors[0].reason


async def test_text_that_looks_like_a_formula_is_exported_as_text(db, other_db, embedder):
    # openpyxl writes any string starting with "=" as a live formula: opening the
    # export would run it, and on re-import it reads back as empty.
    formula_q = '=HYPERLINK("http://example.com","Innova E1 help")'
    category = Category(name="Innova")
    db.add(category)
    await db.commit()
    faq = await add_faq(db, formula_q, E1_A)
    faq.category_id = category.id
    await db.commit()

    exported = await export_knowledge_base(db)
    cell = load_workbook(io.BytesIO(exported))[FAQ_SHEET_NAME]["B2"]
    assert cell.data_type == "s" and cell.value == formula_q

    summary = await import_knowledge_base(other_db, exported)
    assert summary.faqs.created == 1
    assert [f.question for f in (await other_db.execute(select(FAQEntry))).scalars()] == [formula_q]


def test_standalone_faq_export_is_protected_too():
    from app.services.kb_export import text_cells_only

    wb = Workbook()
    wb.active.append(["=1+1", "plain"])
    text_cells_only(wb.active)
    assert [c.data_type for c in wb.active[1]] == ["s", "s"]


async def test_phrasings_that_could_never_be_served_are_refused(db, embedder):
    ni2_q = "Ni2 heater: display dims after an hour, what should we do?"
    await add_faq(db, ni2_q, "On Ni2 heaters the display dims by design after 60 minutes.")
    await add_faq(db, E1_Q, E1_A)
    summary = await import_knowledge_base(db, workbook(phrasing_rows=[
        (ni2_q, "NS heater: display dims after an hour", "TRUE"),  # another heater type
        (E1_Q, "Innova 显示 E1 是什么意思", "TRUE"),  # English answer for a Chinese wording
        (E1_Q, "Innova E1 meaning, answer in Chinese", "TRUE"),  # asks for a reply language
    ]))
    assert summary.phrasings.failed == 3 and summary.phrasings.created == 0
    assert all("can't be answered with this FAQ" in e.reason for e in summary.phrasings.errors)
